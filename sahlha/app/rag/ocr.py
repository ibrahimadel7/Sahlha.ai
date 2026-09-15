"""OCR abstraction. `extract_document_text` distinguishes text vs scanned docs.

Provider can be swapped later; rest of the app only depends on this interface.
"""
from __future__ import annotations

import io
from dataclasses import dataclass


@dataclass
class ExtractedDocument:
    text: str
    num_pages: int
    is_scanned: bool  # True when native text extraction yielded ~nothing (image-based)
    method: str       # e.g. "pypdf" | "docx" | "txt" | "ocr:tesseract" | "ocr:unavailable"


def _extract_pdf_native(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return "\n".join(parts), len(reader.pages)


def _extract_docx(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs)


def _tesseract_cmd() -> str | None:
    """Resolve the tesseract binary: PATH, TESSERACT_CMD, then known install locations."""
    import os as _os
    import shutil as _sh

    found = _sh.which("tesseract") or _os.getenv("TESSERACT_CMD")
    if found:
        return found
    for candidate in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                      r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"):
        if _os.path.exists(candidate):
            return candidate
    return None


def _poppler_path() -> str | None:
    """Resolve poppler bin dir (for pdf2image): PATH, POPPLER_PATH, then WinGet location."""
    import glob as _glob
    import os as _os
    import shutil as _sh

    if _sh.which("pdftoppm"):
        return None  # on PATH: pdf2image finds it unaided
    env = _os.getenv("POPPLER_PATH")
    if env and _os.path.exists(_os.path.join(env, "pdftoppm.exe")):
        return env
    base = _os.path.join(_os.getenv("LOCALAPPDATA", ""), "Microsoft", "WinGet", "Packages")
    for hit in _glob.glob(_os.path.join(base, "*oppler*", "**", "pdftoppm.exe"), recursive=True):
        return _os.path.dirname(hit)
    return None


def _try_ocr_images(data: bytes, suffix: str) -> tuple[str, str]:
    """Real OCR via tesseract (native text failed or image input). Returns (text, method)."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return "", "ocr:unavailable (pytesseract/Pillow not installed)"
    tesseract = _tesseract_cmd()
    if not tesseract:
        return "", "ocr:unavailable (tesseract binary not found; set TESSERACT_CMD)"
    pytesseract.pytesseract.tesseract_cmd = tesseract
    try:
        if suffix == ".pdf":
            try:
                from pdf2image import convert_from_bytes
            except ImportError:
                return "", "ocr:unavailable (pdf2image not installed)"
            poppler = _poppler_path()
            if poppler is None and __import__("shutil").which("pdftoppm") is None:
                return "", "ocr:unavailable (poppler not found; set POPPLER_PATH)"
            kwargs = {"poppler_path": poppler} if poppler else {}
            images = convert_from_bytes(data, dpi=200, **kwargs)
            texts = [pytesseract.image_to_string(img) for img in images]
            return "\n".join(texts), "ocr:tesseract(pdf2image)"
        image = Image.open(io.BytesIO(data))
        return pytesseract.image_to_string(image), "ocr:tesseract"
    except Exception as exc:  # OCR is best-effort; never crash ingestion
        return "", f"ocr:failed ({exc})"


def extract_document_text(file_bytes: bytes, filename: str) -> ExtractedDocument:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix == ".pdf":
        text, pages = _extract_pdf_native(file_bytes)
        if text.strip() and len(text.strip()) >= 50:
            return ExtractedDocument(text=text.strip(), num_pages=pages, is_scanned=False, method="pypdf")
        ocr_text, method = _try_ocr_images(file_bytes, ".pdf")
        if ocr_text.strip():
            return ExtractedDocument(text=ocr_text.strip(), num_pages=pages, is_scanned=True, method=method)
        return ExtractedDocument(text=text.strip(), num_pages=pages, is_scanned=True, method=method)
    if suffix == ".docx":
        return ExtractedDocument(text=_extract_docx(file_bytes).strip(), num_pages=1, is_scanned=False, method="docx")
    if suffix in {".txt", ".md", ""}:
        return ExtractedDocument(text=file_bytes.decode("utf-8", errors="ignore").strip(),
                                 num_pages=1, is_scanned=False, method="txt")
    if suffix in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
        ocr_text, method = _try_ocr_images(file_bytes, suffix)
        return ExtractedDocument(text=ocr_text.strip(), num_pages=1, is_scanned=True, method=method)
    # Fallback: try raw decode
    return ExtractedDocument(text=file_bytes.decode("utf-8", errors="ignore").strip(),
                             num_pages=1, is_scanned=False, method="raw")
