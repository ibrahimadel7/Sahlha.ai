"""Golden content fixtures exercise structure, coverage, evidence and grounding."""
import io
import sys
from types import SimpleNamespace

import numpy as np
import pytest

from sahlha.app.agent.agent import SahlhaAgent
from sahlha.app.agent.learning_content import LearningContent
from sahlha.app.agent.llm import fallback_questions
from sahlha.app.agent.tools import content_tools, critique_tools, quality_tools
from sahlha.app.config import settings
from sahlha.app.database.repositories import repositories as repo
from sahlha.app.rag import ocr, ingestion, vectorstore, embeddings
from sahlha.app.rag.chunking import chunk_blocks
from sahlha.app.rag.document import DocumentBlock, text_blocks
from tests.test_upload_edges import text_pdf


GOLDEN = {
    'programming': '# While Loops\nA while loop repeats instructions while a condition is true. The loop stops when the condition is false.\n\n```python\nwhile count < 3:\n    print(count)\n    count += 1\n```\n# Loop Termination\nA break statement ends a loop immediately. An infinite loop never reaches termination.',
    'math': '# Fractions\nFractions are equal parts of a whole. The denominator describes the number of equal parts.\n\nx = 1/2 + 1/4 = 3/4',
    'science': '# Photosynthesis\nPhotosynthesis is the process plants use to convert light into chemical energy. Chlorophyll is a pigment that absorbs light.\n\n| Input | Output |\n| Light | Chemical energy |',
    'history': '# The Industrial Revolution\nThe Industrial Revolution is a period of major manufacturing changes. Mechanization is the use of machines to perform work.\n\n- Factories expanded production.\n- Steam engines powered machinery.',
    'geography': '# Latitude and Longitude\nLatitude is angular distance north or south of the equator. Longitude is angular distance east or west of the prime meridian.',
}


@pytest.mark.parametrize('subject', GOLDEN)
def test_golden_subject_evidence_and_map(db_session, subject):
    ingestion.ingest_upload(db_session, file_bytes=GOLDEN[subject].encode(), filename=f'{subject}.md',
                           course_id='golden', lesson_id=subject, defer_index=True)
    mapped, chunks = content_tools.build_content_map(db_session, 'golden', subject)
    assert sum(len(s['chunk_ids']) for s in mapped['sections']) == len(chunks)
    result = SahlhaAgent(db_session).extract_skills(course_id='golden', lesson_id=subject)
    assert result['skills']
    ids = {c['chunk_id'] for c in chunks}
    for skill in result['skills']:
        assert set(skill['evidence_chunk_ids']) <= ids
        assert skill['evidence_chunk_ids'] and skill['learning_objective']
        assert skill['name'].lower() not in {'false', 'true', 'looping', 'loops', 'output', 'example'}


def test_full_coverage_past_old_top_k(db_session):
    text = '\n\n'.join(f'# Topic {i}\nTopic {i} is an independently defined concept with supporting lesson details.' for i in range(14))
    ingestion.ingest_upload(db_session, file_bytes=text.encode(), filename='long.md', course_id='c', lesson_id='l', defer_index=True)
    result = SahlhaAgent(db_session).extract_skills(course_id='c', lesson_id='l', max_skills=20)
    chunks = repo.get_chunks(db_session, course_id='c', lesson_id='l')
    covered = {i for s in result['skills'] for i in s['evidence_chunk_ids']}
    assert {c.id for c in chunks} <= covered
    assert any('13' in s['name'] for s in result['skills'])


def test_garbage_and_duplicate_skills_rejected():
    chunks = [{'chunk_id': 'a', 'section_id': 's', 'text': 'While loops repeat instructions while a condition is true. False ends the loop.'}]
    candidates = [{'skill_id': name.lower(), 'name': name, 'evidence_chunk_ids': ['a']}
                  for name in ['False', 'Looping', 'Loops', 'While Loop Fundamentals', 'While Loop Basics']]
    good, warnings = content_tools.validate_skills(candidates, chunks, 6)
    assert len(good) == 1 and good[0]['name'] == 'While Loop Fundamentals'
    assert len(warnings) == 3
    assert not content_tools.validate_skills([{'name': 'While Loops', 'evidence_chunk_ids': ['foreign']}], chunks, 6)[0]


def test_code_table_formula_heading_chunk_structure():
    blocks = text_blocks(GOLDEN['programming'])
    code = next(b for b in blocks if b.type == 'code')
    assert '    print(count)' in code.text
    chunks = chunk_blocks(blocks, document_id='d', chunk_size=160, chunk_overlap=20)
    assert any(code.text in c['text'] for c in chunks)
    assert all(c['page'] == 1 and c['section_id'] for c in chunks)
    assert chunks[0]['text'].startswith('# While Loops\n\nA while loop')
    assert any(b.type == 'table' for b in text_blocks(GOLDEN['science']))
    assert any(b.type == 'formula' for b in text_blocks(GOLDEN['math']))
    oversized = chunk_blocks([DocumentBlock(2, 0, 'code', 'x' * 2000)], chunk_size=100, chunk_overlap=10)
    assert all(len(c['text']) <= 400 for c in oversized)
    assert ''.join(c['text'] for c in oversized) == 'x' * 2000
    assert oversized[-1]['block_metadata']['continuation']


def test_docx_table_reading_order():
    from docx import Document
    doc = Document()
    doc.add_heading('Cell Structure', 1)
    doc.add_paragraph('Cells contain specialized structures that perform essential functions.')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Nucleus'
    table.cell(0, 1).text = 'Stores DNA'
    doc.add_paragraph('Review the organelles.', style='List Bullet')
    data = io.BytesIO()
    doc.save(data)
    result = ocr.extract_document_text(data.getvalue(), 'cells.docx')
    assert [b.type for b in result.blocks] == ['heading', 'paragraph', 'table', 'list']
    assert result.blocks[2].text == 'Nucleus\tStores DNA'
    assert result.blocks[0].metadata['page_is_logical']


def test_pptx_slide_boundaries():
    from pptx import Presentation
    presentation = Presentation()
    for title in ('Equations', 'Graphs'):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = f'{title} are representations used in mathematics.'
    data = io.BytesIO()
    presentation.save(data)
    result = ocr.extract_document_text(data.getvalue(), 'math.pptx')
    assert result.num_pages == 2
    assert [b.page for b in result.blocks] == [1, 1, 2, 2]
    assert result.blocks[0].type == 'heading'


def test_mixed_pdf_only_ocr_sparse_page(monkeypatch):
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    for text in ('A clean native lesson explains fractions as equal parts of a whole.', 'Page 2'):
        writer.add_page(PdfReader(io.BytesIO(text_pdf(text))).pages[0])
    data = io.BytesIO()
    writer.write(data)
    calls = []
    def scan(data, suffix):
        calls.append(len(PdfReader(io.BytesIO(data)).pages))
        return 'Recovered page explains denominators and numerators.', 'ocr:tesseract'
    monkeypatch.setattr(ocr, '_try_ocr_images', scan)
    result = ocr.extract_document_text(data.getvalue(), 'mixed.pdf')
    assert calls == [1] and result.num_pages == 2
    assert {b.page for b in result.blocks} == {1, 2}
    assert result.method == 'hybrid:pypdf+tesseract'
    assert 'clean native' in result.text and 'Recovered' in result.text


def test_broken_ocr_warning(monkeypatch):
    monkeypatch.setattr(ocr, '_try_ocr_images', lambda *a: ('\ufffd' * 100, 'ocr:tesseract'))
    result = ocr.extract_document_text(b'scan', 'scan.png')
    assert result.warnings and result.quality['quality_indicator'] == 0


@pytest.mark.parametrize('dimension', [256, 768, 1024])
def test_dynamic_embedding_dimensions(monkeypatch, dimension):
    class Model:
        def __init__(self, name): pass
        def get_sentence_embedding_dimension(self): return dimension
        def encode(self, texts, **kwargs): return np.ones((len(texts), dimension)) / dimension ** .5
    monkeypatch.setitem(sys.modules, 'sentence_transformers', SimpleNamespace(SentenceTransformer=Model))
    model = embeddings.DenseEmbeddingModel('test')
    assert model.embed(['lesson']).shape == (1, dimension)


def test_hybrid_scope_evidence_links_and_reranking(db_session, monkeypatch):
    from tests.test_rag_upgrade import DenseFake
    monkeypatch.setattr(embeddings, '_embeddings', DenseFake())
    for lesson in ['a', 'b']:
        ingestion.ingest_upload(db_session, file_bytes=b'Loops repeat instructions. Iteration repeats a sequence.',
            filename='lesson.txt', course_id='c', lesson_id=lesson, defer_index=True)
    chunks = repo.get_chunks(db_session, course_id='c', lesson_id='a')
    repo.upsert_skill(db_session, course_id='c', lesson_id='a', skill_id='loop',
        educational_metadata={'evidence_chunk_ids': [chunks[0].id]})
    found = vectorstore.search(db_session, 'iteration', course_id='c', lesson_id='a', skill_id='loop')
    assert found and all(c['lesson_id'] == 'a' for c in found)
    assert found[0]['retrieval_source'] == 'hybrid'
    assert not vectorstore.search(db_session, 'iteration', course_id='c', lesson_id='', skill_id='loop')
    monkeypatch.setattr(settings, 'reranker_enabled', True)
    monkeypatch.setitem(vectorstore._rerankers, settings.reranker_model, SimpleNamespace(predict=lambda pairs: np.arange(len(pairs))))
    assert vectorstore.rerank('x', [SimpleNamespace(text='one'), SimpleNamespace(text='two')], [0, 1]) == [1, 0]
    monkeypatch.setitem(vectorstore._rerankers, settings.reranker_model, SimpleNamespace(predict=lambda pairs: [float('nan')] * len(pairs)))
    assert vectorstore.rerank('x', [SimpleNamespace(text='one')], [0]) == [0]


def test_question_grounding_detects_wrong_answer_and_foreign_citation():
    context = [{'chunk_id': 'a', 'text': 'Chlorophyll absorbs light energy in green plants. Photosynthesis produces chemical energy from sunlight.'}]
    questions = fallback_questions(context, 'photosynthesis', 4)
    assert len(questions) == 4
    for q in questions:
        assert critique_tools.critique_question(q, context, 'photosynthesis')[0]
        assert not critique_tools.critique_question(dict(q, correct_answer=(q['correct_answer']+1)%4), context, 'photosynthesis')[0]
        assert not critique_tools.critique_question(dict(q, evidence_chunk_ids=['foreign']), context, 'photosynthesis')[0]
    with pytest.raises(ValueError, match='Not enough grounded'):
        critique_tools.critique_and_top_up([], [{'chunk_id': 'x', 'text': '???'}], 'x', 4)


def test_safe_learning_visual_schema():
    content = LearningContent(core_idea='Track loop state.', visual_type='code_trace',
        visual_spec={'source_text': 'while x < 2:', 'items': [{'label': 'x', 'value': 1}]},
        playground={'interaction': 'step'})
    assert content.visual_type == 'code_trace'
    with pytest.raises(ValueError):
        LearningContent(core_idea='x', visual_type='html', visual_spec={'script': 'alert(1)'})
    with pytest.raises(ValueError):
        LearningContent(core_idea='x', visual_type='code_trace', playground={'execute': 'anything'})


def test_quality_signals_not_yet_generated(db_session):
    ingestion.ingest_upload(db_session, file_bytes=GOLDEN['science'].encode(), filename='science.md',
        course_id='c', lesson_id='l', defer_index=True)
    quality = quality_tools.lesson_quality(db_session, 'c', 'l')
    assert quality['extraction_quality'] > .65
    assert quality['question_grounding_quality'] is None


def test_legacy_skill_refresh_retains_history(db_session):
    ingestion.ingest_upload(db_session, file_bytes=GOLDEN['programming'].encode(), filename='loops.md',
        course_id='c', lesson_id='l', defer_index=True)
    old = repo.upsert_skill(db_session, course_id='c', lesson_id='l', skill_id='false', name='False')
    old_id = old.id
    result = SahlhaAgent(db_session).extract_skills(course_id='c', lesson_id='l')
    assert all(s['name'] != 'False' for s in result['skills'])
    assert repo.get_skill(db_session, course_id='c', lesson_id='l', skill_id='false').id == old_id
    assert not old.extraction_active
    assert old not in repo.list_skills(db_session, course_id='c', lesson_id='l')
    again = SahlhaAgent(db_session).extract_skills(course_id='c', lesson_id='l')
    assert again['backend'] == 'existing'


def test_question_evidence_and_verification_persist(db_session):
    ingestion.ingest_upload(db_session, file_bytes=GOLDEN['science'].encode(), filename='science.md',
        course_id='c', lesson_id='l', skill_id='science', defer_index=True)
    bank = SahlhaAgent(db_session).generate_question_bank(course_id='c', lesson_id='l', skill_id='science', n_questions=4)
    questions = repo.get_questions(db_session, bank['question_bank_id'])
    assert len({q.question_text for q in questions}) == 4
    assert all(q.evidence_chunk_ids and q.verification['passed'] and q.learning_objective for q in questions)
    assert quality_tools.lesson_quality(db_session, 'c', 'l')['question_grounding_quality'] == 1


def test_optional_verifier_requires_all_checks(monkeypatch):
    monkeypatch.setattr(settings, 'enable_llm_critique', True)
    context = [{'chunk_id': 'a', 'text': 'Chlorophyll absorbs light energy in green plants. Photosynthesis produces chemical energy from sunlight.'}]
    question = {'question': 'What absorbs light energy in green plants?', 'type': 'multiple_choice',
        'options': ['Chlorophyll', 'Hemoglobin', 'Insulin', 'Keratin'], 'correct_answer': 0,
        'difficulty': 'easy', 'evidence_chunk_ids': ['a']}
    monkeypatch.setattr(critique_tools, 'complete_json', lambda *a, **k: ({'valid': True}, 'fake'))
    result, meta = critique_tools.critique_and_top_up([question], context, 'science', 1)
    assert meta['rejected'] and result[0]['verification']['method'] == 'source_completion'
    checks = {k: True for k in ['answerable', 'answer_supported', 'distractors_incorrect', 'unambiguous', 'clear', 'difficulty', 'objective']}
    monkeypatch.setattr(critique_tools, 'complete_json', lambda *a, **k: ({'valid': True, 'checks': checks}, 'fake'))
    result, meta = critique_tools.critique_and_top_up([question], context, 'science', 1)
    assert not meta['rejected'] and result[0]['verification']['method'] in ('llm', 'llm_verified')


def test_docx_math_preserved():
    from docx import Document
    from docx.oxml import parse_xml
    document = Document()
    paragraph = document.add_paragraph()
    paragraph._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x = 2</m:t></m:r></m:oMath>'))
    data = io.BytesIO()
    document.save(data)
    extracted = ocr.extract_document_text(data.getvalue(), 'math.docx')
    assert extracted.blocks[0].text == 'x = 2' and extracted.blocks[0].type == 'formula'


def test_additive_content_migration(tmp_path, monkeypatch):
    from sqlalchemy import create_engine, text
    from sqlalchemy.orm import Session
    from sahlha.app.database import database
    engine = create_engine(f'sqlite:///{tmp_path / "legacy-content.db"}')
    with engine.begin() as connection:
        connection.execute(text('''CREATE TABLE skills (id VARCHAR(32) PRIMARY KEY,
            course_id VARCHAR(128), lesson_id VARCHAR(128), skill_id VARCHAR(128),
            name VARCHAR(256), description TEXT, explanation TEXT, key_concepts JSON,
            image_url VARCHAR(1024), image_path VARCHAR(1024), image_alt VARCHAR(512),
            audio_path VARCHAR(1024), created_at DATETIME, updated_at DATETIME)'''))
        connection.execute(text("INSERT INTO skills (id, course_id, lesson_id, skill_id, name, explanation) VALUES ('old', 'c', 'l', 's', 'Existing topic', 'Keep this explanation')"))
    monkeypatch.setattr(database, 'engine', engine)
    database.init_db()
    database.init_db()
    with Session(engine) as db:
        row = repo.get_skill(db, course_id='c', lesson_id='l', skill_id='s')
        assert row.id == 'old' and row.explanation == 'Keep this explanation'
        assert row.evidence_chunk_ids == [] and row.learning_content == {} and row.extraction_active
    engine.dispose()


def test_content_map_invalidated_when_location_changes(db_session):
    ingestion.ingest_upload(db_session, file_bytes=GOLDEN['math'].encode(), filename='math.md',
        course_id='c', lesson_id='l', defer_index=True)
    before, _ = content_tools.build_content_map(db_session, 'c', 'l')
    chunk = repo.get_chunks(db_session, course_id='c', lesson_id='l')[0]
    chunk.page = 3
    db_session.commit()
    after, _ = content_tools.build_content_map(db_session, 'c', 'l')
    assert before['source_fingerprint'] != after['source_fingerprint']
    assert after['sections'][0]['page'] == 3


def test_ocr_preprocessing_language_and_orientation(monkeypatch):
    from PIL import Image
    import pytesseract
    data = io.BytesIO()
    Image.new('RGB', (80, 40), 'white').save(data, format='PNG')
    monkeypatch.setattr(settings, 'ocr_languages', 'ara+eng')
    monkeypatch.setattr(pytesseract, 'image_to_osd', lambda *a, **k: {'orientation_conf': 12, 'rotate': 90})
    def recognize(image, **kwargs):
        assert image.mode == 'L' and image.size == (40, 80)
        assert kwargs['lang'] == 'ara+eng' and 'preserve_interword_spaces=1' in kwargs['config']
        return 'Recognized lesson source text from an oriented image.'
    monkeypatch.setattr(pytesseract, 'image_to_string', recognize)
    extracted = ocr.extract_document_text(data.getvalue(), 'image.png')
    assert extracted.quality['ocr_used'] and not extracted.warnings


def test_prompt_packs_complete_code_blocks():
    from sahlha.app.agent.prompts import evidence_context
    text = 'while count < 3:\n    print(count)\n    count += 1'
    chunks = [{'chunk_id': 'a', 'text': text}, {'chunk_id': 'b', 'text': 'x' * 500}]
    context = evidence_context(chunks, budget=150)
    assert text in context and '[chunk b' not in context
