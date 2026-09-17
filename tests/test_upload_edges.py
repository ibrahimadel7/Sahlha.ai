"""Regression coverage for document upload failures and extraction."""
import io

from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from sahlha.app.rag.ocr import extract_document_text
from tests.test_platform import _setup_teacher_student, _auth


def text_pdf(text="Fractions are parts of a whole."):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
        DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    result = io.BytesIO()
    writer.write(result)
    return result.getvalue()


def test_short_pdf_does_not_require_ocr(monkeypatch):
    def unexpected(*args):
        raise AssertionError("Native text must not require OCR")
    monkeypatch.setattr("sahlha.app.rag.ocr._try_ocr_images", unexpected)
    result = extract_document_text(text_pdf(), "lesson.pdf")
    assert result.method == "pypdf"
    assert "Fractions" in result.text


def test_pptx_and_docx_tables():
    from pptx import Presentation
    from pptx.util import Inches
    from docx import Document
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "Fraction lesson"
    slide.shapes.add_table(1, 1, 0, Inches(1), Inches(3), Inches(1)).table.cell(0, 0).text = "Numerator"
    data = io.BytesIO()
    ppt.save(data)
    assert "Numerator" in extract_document_text(data.getvalue(), "lesson.pptx").text
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Denominator"
    data = io.BytesIO()
    doc.save(data)
    assert "Denominator" in extract_document_text(data.getvalue(), "lesson.docx").text


def test_pdf_upload_success_and_corruption(client, monkeypatch, tmp_path):
    monkeypatch.setattr("sahlha.app.config.settings.upload_dir", str(tmp_path))
    teacher, _, room = _setup_teacher_student(client)
    def upload(data):
        return client.post("/materials/upload", headers=_auth(teacher["token"]),
            data={"classroom_id": room["id"]}, files={"file": ("lesson.pdf", data, "application/pdf")})
    result = upload(text_pdf())
    assert result.status_code == 201
    assert result.json()["material"]["status"] == "processing"
    material_id = result.json()["material"]["id"]
    assert client.get(f"/materials/{material_id}", headers=_auth(teacher["token"])).json()["status"] == "processed"
    result = upload(b"broken pdf")
    assert result.status_code == 201
    assert result.json()["material"]["status"] == "failed"
    assert result.json()["error"]
    assert upload(text_pdf()).json()["material"]["status"] == "processing"


def test_scan_missing_ocr_has_actionable_error(client, monkeypatch, tmp_path):
    monkeypatch.setattr("sahlha.app.config.settings.upload_dir", str(tmp_path))
    monkeypatch.setattr("sahlha.app.rag.ocr._try_ocr_images", lambda *args: ("", "ocr:unavailable"))
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    data = io.BytesIO()
    writer.write(data)
    teacher, _, room = _setup_teacher_student(client)
    result = client.post("/materials/upload", headers=_auth(teacher["token"]),
        data={"classroom_id": room["id"]}, files={"file": ("scan.pdf", data.getvalue())})
    assert result.json()["material"]["status"] == "failed"
    assert "OCR" in result.json()["error"]


def test_scanned_pdf_is_rendered_one_page_at_a_time(monkeypatch):
    import pdf2image
    import pytesseract
    from sahlha.app.rag.ocr import _try_ocr_images
    rendered = []
    closed = []
    class Image:
        def __init__(self, page):
            self.page = page
        def close(self):
            closed.append(self.page)
    def render(data, **kwargs):
        assert kwargs["first_page"] == kwargs["last_page"]
        assert 0 < kwargs["timeout"] <= 45
        rendered.append(kwargs["first_page"])
        return [Image(kwargs["first_page"])]
    monkeypatch.setattr(pdf2image, "pdfinfo_from_bytes", lambda *args, **kwargs: {"Pages": 3})
    monkeypatch.setattr(pdf2image, "convert_from_bytes", render)
    monkeypatch.setattr(pytesseract, "image_to_string", lambda image, **kwargs: f"Page {image.page}")
    out = _try_ocr_images(b"scan", ".pdf")
    text, method = out[0], out[1]
    assert rendered == closed == [1, 2, 3]
    assert "Page 3" in text
    assert method == "ocr:tesseract(pdf2image)"


def test_scan_timeout_does_not_report_partial_success(monkeypatch):
    import pdf2image
    import pytesseract
    from sahlha.app.rag.ocr import _try_ocr_images
    monkeypatch.setattr(pdf2image, "pdfinfo_from_bytes", lambda *args, **kwargs: {"Pages": 1})
    monkeypatch.setattr(pdf2image, "convert_from_bytes", lambda *args, **kwargs: [])
    times = iter([0, 46])
    monkeypatch.setattr("sahlha.app.rag.ocr.time.monotonic", lambda: next(times))
    out = _try_ocr_images(b"scan", ".pdf")
    text, method = out[0], out[1]
    assert text == ""
    assert method.startswith("ocr:failed")
    assert "smaller scan" in method
